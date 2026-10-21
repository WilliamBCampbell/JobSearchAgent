import json
import subprocess
from pathlib import Path
import pytest
from docx import Document
from jobsearch.memory import Memory, FactInput
from jobsearch.corpus import Corpus
from jobsearch.config import atomic_write
from jobsearch.providers import Providers, ProviderConfig


def fact(db, key="citizenship", value="US", scope="global"):
    return Memory(db).put(
        FactInput(
            key=key,
            value=value,
            status="confirmed",
            scope=scope,
            provenance={"type": "user_confirmation", "source": "synthetic test"},
        )
    )


def test_repeated_question_uses_memory_without_rag(db):
    f = fact(db)
    memory = Memory(db)
    for _ in range(3):
        assert memory.resolve("citizenship")["value"] == "US"
    assert db.one("SELECT value FROM metrics WHERE name='memory_hits'")["value"] == 3
    assert not db.one("SELECT 1 FROM metrics WHERE name='retrieval_calls'")


def test_correction_invalidates_cached_answer(db):
    old = fact(db, value="old")
    memory = Memory(db)
    memory.resolve("citizenship")
    memory.put(
        FactInput(
            key="citizenship",
            value="corrected",
            status="confirmed",
            supersedes=old["id"],
            provenance={"type": "user_confirmation"},
        )
    )
    assert db.one("SELECT valid FROM answer_cache")["valid"] is False
    assert memory.resolve("citizenship")["value"] == "corrected"
    assert (
        db.one("SELECT status FROM facts WHERE id=%s", (old["id"],))["status"]
        == "superseded"
    )


@pytest.mark.parametrize(
    "key",
    [
        "salary",
        "travel",
        "relocation",
        "schedule",
        "reference_approval",
        "submission_authorization",
    ],
)
def test_role_answers_do_not_leak(db, key):
    fact(db, key, "role A only", "role-a")
    assert Memory(db).resolve(key, scope="role-a")["value"] == "role A only"
    assert Memory(db).resolve(key, scope="role-b")["source"] == "question"
    with pytest.raises(ValueError):
        fact(db, key, "global leak")


def test_conflict_requires_resolution(db):
    fact(db, value="one")
    fact(db, value="two")
    answer = Memory(db).resolve("citizenship")
    assert answer["source"] == "question" and len(answer["conflicting_facts"]) == 2


@pytest.mark.parametrize("key", ["password", "verification_code", "session_token"])
def test_secrets_are_never_memory(db, key):
    with pytest.raises(ValueError):
        fact(db, key, "not-to-be-stored")


def test_generated_draft_cannot_confirm_a_claim(db):
    path = db.settings.fixture_workspace / "Generated_Resume.md"
    path.write_text("I hold an unsupported doctorate.")
    result = Corpus(db).ingest(path.name, "fixture", kind="generated_draft")
    with pytest.raises(ValueError):
        Memory(db).put(
            FactInput(
                key="degree",
                value="doctorate",
                status="confirmed",
                provenance={"type": "explicit_record"},
                evidence_versions=[result["version"]],
            ),
            actor="importer",
        )
    assert Memory(db).resolve("degree")["source"] == "question"


def test_changed_and_deleted_evidence_invalidate_memory_and_cache(db):
    path = db.settings.fixture_workspace / "Confirmed_Answers.json"
    path.write_text(json.dumps({"confirmed": {"skill": "C#"}}))
    corpus = Corpus(db)
    v = corpus.ingest(path.name, "fixture", kind="confirmation_record")
    Memory(db).put(
        FactInput(
            key="skill",
            value="C#",
            status="confirmed",
            provenance={"type": "explicit_record"},
            evidence_versions=[v["version"]],
        ),
        actor="importer",
    )
    Memory(db).resolve("skill")
    path.write_text(json.dumps({"confirmed": {"skill": "Corrected"}}))
    corpus.ingest(path.name, "fixture")
    assert db.one("SELECT active FROM facts")["active"] is False
    assert db.one("SELECT valid FROM answer_cache")["valid"] is False
    path.unlink()
    corpus.sync("fixture")
    assert db.one("SELECT deleted FROM documents")["deleted"] is True
    assert len(db.all("SELECT * FROM versions")) == 2


def test_duplicate_formats_do_not_dominate(db):
    root = db.settings.fixture_workspace
    (root / "Package.md").write_text("# API delivery\nC# APIs and database work.")
    (root / "Package.html").write_text(
        "<h1>API delivery</h1><p>C# APIs and database work.</p>"
    )
    Corpus(db).sync("fixture")
    result = Corpus(db).retrieve("C# APIs", "fixture")
    assert len(result["evidence"]) == 1


def test_later_pdf_pages_and_full_document_are_retrieved(db):
    root = db.settings.fixture_workspace
    doc = Document()
    doc.add_paragraph("An ordinary career overview.")
    doc.add_page_break()
    doc.add_paragraph(
        "Project detail: implemented a production data sanitization pipeline."
    )
    doc.add_page_break()
    doc.add_paragraph(
        "The Zephyr integration generated traceable test cases from support tickets."
    )
    source = root / "Career.docx"
    doc.save(source)
    subprocess.run(
        [
            "libreoffice",
            f"-env:UserInstallation=file://{root}/lo-test",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(root),
            str(source),
        ],
        check=True,
        capture_output=True,
        timeout=60,
    )
    Corpus(db).ingest("Career.pdf", "fixture")
    result = Corpus(db).retrieve(
        "Zephyr traceable test cases", "fixture", full_document=True
    )
    hit = result["evidence"][0]
    assert any(
        c["locator"].get("page") == 3 and "Zephyr" in c["body"] for c in hit["context"]
    )
    assert {c["locator"]["page"] for c in hit["context"]} == {1, 2, 3}
    assert hit["expansion"] == "full_document"


def test_ocr_failure_is_visible(db):
    from pypdf import PdfWriter

    p = db.settings.fixture_workspace / "Scanned.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=600, height=800)
    writer.write(p)
    result = Corpus(db).ingest(p.name, "fixture")
    assert result["status"] == "failed" and result["warnings"]


def test_provider_disabled_prevents_network(db, monkeypatch):
    provider = Providers(db)
    provider.configure(
        ProviderConfig(
            provider="ollama",
            destination="https://example.com",
            enabled=False,
            embedding_model="test",
        )
    )
    called = []
    monkeypatch.setattr("httpx.Client.post", lambda *a, **k: called.append(1))
    with pytest.raises(ValueError, match="Enable"):
        provider.embed(["private document"])
    assert not called


def test_expired_fact_is_not_answered_from_cache(db):
    f = fact(db)
    Memory(db).resolve("citizenship")
    db.execute("UPDATE facts SET effective_to='2000-01-01' WHERE id=%s", (f["id"],))
    assert Memory(db).resolve("citizenship")["source"] == "question"
    assert (
        db.one("SELECT value FROM metrics WHERE name='stale_answers_blocked'")["value"]
        == 1
    )


def test_cache_key_includes_fact_and_document_versions(db):
    old = fact(db)
    Memory(db).resolve("citizenship")
    first = db.one("SELECT key,request_key FROM answer_cache")
    Memory(db).put(
        FactInput(
            key="citizenship",
            value="new",
            status="confirmed",
            supersedes=old["id"],
            provenance={"type": "user_confirmation"},
        )
    )
    Memory(db).resolve("citizenship")
    rows = db.all("SELECT key,request_key FROM answer_cache")
    assert len(rows) == 2 and len({r["key"] for r in rows}) == 2
    assert {r["request_key"] for r in rows} == {first["request_key"]}


def test_ods_retrieval_retains_header_and_row(db):
    from jobsearch.fixtures import make_tracker
    from jobsearch.tracker import TRACKER_NAME

    make_tracker(db.settings.fixture_workspace / TRACKER_NAME)
    Corpus(db).ingest(TRACKER_NAME, "fixture")
    hits = Corpus(db).retrieve("Deferred Archive Demo", "fixture")["evidence"]
    contexts = [c for h in hits for c in h["context"]]
    assert any(
        c["locator"].get("row") == 7
        and "Company | Role" in c["body"]
        and "AD-002" in c["body"]
        for c in contexts
    )


def test_explicit_historical_sections_are_excluded(db):
    path = db.settings.fixture_workspace / "Notes.md"
    path.write_text(
        "# Current\nC# APIs.\n\n# Superseded history\nUnsupported primary framework ownership.\n"
    )
    Corpus(db).ingest(path.name, "fixture")
    result = Corpus(db).retrieve("framework ownership", "fixture", full_document=True)
    assert all(
        "Unsupported" not in c["body"]
        for hit in result["evidence"]
        for c in hit["context"]
    )
